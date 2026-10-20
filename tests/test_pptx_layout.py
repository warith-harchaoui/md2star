"""Tests for md2star.pptx_layout — catalog extraction, segmentation, selection.

Exercised **offline**: the deterministic catalog/segment stages need only
``python-pptx`` (its own bundled default template is a real, always-available
PPTX with no external asset), and the LLM/VLM stages are driven with injected
fakes (mirroring tests/test_reverse_twin.py), no Ollama required.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from md2star import pptx_layout as pl  # noqa: E402


@pytest.fixture
def default_template(tmp_path: Path) -> Path:
    """python-pptx's own bundled default template: 11 real layouts, no example slides."""
    path = tmp_path / "default.pptx"
    pptx.Presentation().save(str(path))
    return path


# ── deterministic catalog extraction ─────────────────────────────────────────
class TestExtractLayouts:
    def test_extracts_every_layout_even_with_no_example_slides(
        self, default_template: Path
    ) -> None:
        layouts = pl.extract_layouts(default_template)
        names = {layout.name for layout in layouts}
        assert "Title Slide" in names
        assert "Picture with Caption" in names
        # No slide in a fresh Presentation() uses any layout yet.
        assert all(layout.example_pages == [] for layout in layouts)

    def test_picture_placeholder_has_geometry(self, default_template: Path) -> None:
        layouts = {layout.name: layout for layout in pl.extract_layouts(default_template)}
        pic = layouts["Picture with Caption"]
        assert pic.has_kind("pic")
        placeholder = next(p for p in pic.placeholders if p.kind == "pic")
        assert placeholder.emu is not None
        assert all(v > 0 for v in placeholder.emu)

    def test_no_name_collapses_into_reference_hints(self, default_template: Path) -> None:
        # None of python-pptx's stock layouts look like brand-reference pages.
        layouts = pl.extract_layouts(default_template)
        assert all(layout.is_content_layout for layout in layouts)

    def test_rep_page_is_none_without_examples(self, default_template: Path) -> None:
        layouts = pl.extract_layouts(default_template)
        assert all(layout.rep_page is None for layout in layouts)


# ── deck segmentation ────────────────────────────────────────────────────────
class TestSegment:
    def test_splits_on_horizontal_rules(self) -> None:
        md = "# Title\n\nIntro\n\n***\n\n## Second\n\n- a\n- b\n"
        slides = pl.segment(md)
        assert len(slides) == 2
        assert slides[0].title == "Title"
        assert slides[1].title == "Second"
        assert "bullets" in slides[1].features

    def test_falls_back_to_headings_without_hr(self) -> None:
        md = "# One\n\nbody\n\n## Two\n\nmore body\n"
        slides = pl.segment(md)
        assert [s.title for s in slides] == ["One", "Two"]

    def test_strips_html_comments_before_titling(self) -> None:
        md = "<!-- note -->\n# Real Title\n\nbody\n"
        slides = pl.segment(md)
        assert slides[0].title == "Real Title"

    def test_detects_features(self) -> None:
        md = (
            "# T\n\n```mermaid\nflowchart LR\nA-->B\n```\n\n"
            "1. step\n\n> a quote\n\n| a | b |\n|---|---|\n"
        )
        slides = pl.segment(md)
        f = set(slides[0].features)
        assert {"mermaid", "ordered-list", "quote", "table"} <= f

    def test_empty_chunks_are_dropped(self) -> None:
        md = "# T\n\nbody\n\n***\n\n   \n\n***\n\n## U\n\nbody2\n"
        slides = pl.segment(md)
        assert [s.title for s in slides] == ["T", "U"]


# ── front matter ─────────────────────────────────────────────────────────────
# A metadata block's closing `---` is a horizontal rule to every regex in the
# module, so left in place it becomes a leading slide whose bullets are the
# author's own metadata. Pandoc consumes it; so must this path.
class TestFrontMatter:
    MD = "---\ntitle: Quarterly Review\nauthor: Jane Roe\n---\n\n# Quarterly Review\n\nbody\n"

    def test_metadata_does_not_become_a_slide(self) -> None:
        slides = pl.segment(self.MD)
        assert [s.title for s in slides] == ["Quarterly Review"]
        assert "author: Jane Roe" not in slides[0].body

    def test_scalars_are_parsed(self) -> None:
        _, meta = pl.split_front_matter(self.MD)
        assert meta == {"title": "Quarterly Review", "author": "Jane Roe"}

    def test_closing_dots_are_accepted(self) -> None:
        """Pandoc closes a metadata block with either `---` or `...`."""
        body, meta = pl.split_front_matter("---\ntitle: T\n...\n\n# H\n")
        assert meta["title"] == "T"
        assert body.strip() == "# H"

    def test_a_metadata_title_names_an_untitled_cover(self) -> None:
        slides = pl.segment("---\ntitle: From Metadata\n---\n\nJust a paragraph.\n")
        assert slides[0].title == "From Metadata"

    def test_only_a_leading_block_counts(self) -> None:
        """Mid-document `---` is a slide break, not metadata."""
        slides = pl.segment("# A\n\nbody\n\n---\n\n# B\n\nbody\n")
        assert [s.title for s in slides] == ["A", "B"]

    def test_a_document_without_metadata_is_untouched(self) -> None:
        body, meta = pl.split_front_matter("# H\n\nbody\n")
        assert meta == {}
        assert body == "# H\n\nbody\n"

    def test_a_deck_that_opens_with_a_slide_break_keeps_its_first_slide(self) -> None:
        """`---` opens a metadata block AND separates slides. Only the content
        inside tells them apart, and guessing wrong deletes a slide."""
        slides = pl.segment("---\n\n# First\n\nbody\n\n---\n\n# Second\n\nbody\n")
        assert [s.title for s in slides] == ["First", "Second"]

    def test_a_key_with_a_nested_value_still_counts_as_metadata(self) -> None:
        body, meta = pl.split_front_matter("---\nauthor:\n  - Jane Roe\ntitle: T\n---\n\n# H\n")
        assert meta["title"] == "T"
        assert body.strip() == "# H"


# ── speaker notes ────────────────────────────────────────────────────────────
# Notes are the one part of a slide that must NOT be drawn on it: the body is
# what the audience reads. Leaving a `::: notes` div in `body` published the
# presenter's private script to the room.
class TestNotes:
    def test_notes_are_lifted_off_the_body(self) -> None:
        slides = pl.segment("## T\n\n- visible\n\n::: notes\nthe private script\n:::\n")
        assert slides[0].notes == "the private script"
        assert "private script" not in slides[0].body

    def test_attribute_form_and_longer_fences(self) -> None:
        slides = pl.segment("## T\n\nvisible\n\n:::: {.notes}\nsecret\n::::\n")
        assert slides[0].notes == "secret"

    def test_several_divs_on_one_slide_are_joined(self) -> None:
        md = "## T\n\nvisible\n\n::: notes\nfirst\n:::\n\nmore\n\n::: notes\nsecond\n:::\n"
        slides = pl.segment(md)
        assert slides[0].notes == "first\n\nsecond"
        assert "more" in slides[0].body

    def test_a_slide_without_notes_has_none(self) -> None:
        assert pl.segment("## T\n\nbody\n")[0].notes == ""

    def test_a_columns_div_is_not_mistaken_for_notes(self) -> None:
        slides = pl.segment("## T\n\n::: columns\nleft\n:::\n")
        assert slides[0].notes == ""
        assert "left" in slides[0].body


# ── catalog build (deterministic-only, no vision) ────────────────────────────
class TestBuildCatalog:
    def test_no_vision_uses_deterministic_extraction(
        self, tmp_path: Path, default_template: Path
    ) -> None:
        cache = tmp_path / "catalog.json"
        catalog = pl.build_catalog(default_template, use_vision=False, cache_path=cache)
        assert {c.name for c in catalog} == {c.name for c in pl.extract_layouts(default_template)}
        assert cache.exists()

    def test_cache_round_trips(self, tmp_path: Path, default_template: Path) -> None:
        cache = tmp_path / "catalog.json"
        first = pl.build_catalog(default_template, use_vision=False, cache_path=cache)
        second = pl.build_catalog(default_template, use_vision=False, cache_path=cache)
        assert [c.name for c in first] == [c.name for c in second]

    def test_vision_captions_when_examples_and_transport_available(self, tmp_path: Path) -> None:
        # Build a template with ONE slide (so the "Title Slide" layout has a
        # renderable example) and a fake render+chat pair.
        prs = pptx.Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        template = tmp_path / "one_slide.pptx"
        prs.save(str(template))

        def fake_render(_template: Path, page: int) -> bytes | None:
            return b"png-bytes" if page == 1 else None

        def fake_chat(_prompt: str, images, _schema):
            assert images  # the caption stage must pass the rendered thumbnail(s)
            return {"archetype": "cover", "is_content_layout": True, "one_line": "a cover slide"}

        cache = tmp_path / "catalog.json"
        catalog = pl.build_catalog(template, chat=fake_chat, render=fake_render, cache_path=cache)
        title_slide = next(c for c in catalog if c.name == "Title Slide")
        assert title_slide.archetype == "cover"
        assert title_slide.caption == "a cover slide"


# ── LLM text-based selection ─────────────────────────────────────────────────
class TestSelectLayouts:
    def test_assigns_catalog_name_from_fake_llm(self) -> None:
        slides = pl.segment("# Cover\n\nsubtitle\n\n***\n\n## Bullets\n\n- a\n- b\n")
        catalog = [
            pl.LayoutInfo(name="Cover Layout", archetype="cover", is_content_layout=True),
            pl.LayoutInfo(
                name="Bulleted Layout", archetype="bulleted-content", is_content_layout=True
            ),
        ]

        def fake_chat(_prompt: str, _images, _schema):
            return {
                "choices": [
                    {"index": 1, "layout_name": "Cover Layout", "confidence": 0.9, "why": "opener"},
                    {
                        "index": 2,
                        "layout_name": "Bulleted Layout",
                        "confidence": 0.8,
                        "why": "list",
                    },
                ]
            }

        pl.select_layouts(slides, catalog, chat=fake_chat)
        assert slides[0].chosen_layout == "Cover Layout"
        assert slides[1].chosen_layout == "Bulleted Layout"
        assert slides[1].chosen_why == "list"

    def test_falls_back_on_unparseable_reply(self) -> None:
        slides = pl.segment("# Cover\n\nbody\n")
        catalog = [
            pl.LayoutInfo(
                name="Bulleted Layout", archetype="bulleted-content", is_content_layout=True
            )
        ]
        pl.select_layouts(slides, catalog, chat=lambda *_: "not json")
        assert slides[0].chosen_layout == "Bulleted Layout"  # never left unassigned

    def test_drops_reference_layouts_from_the_offer(self) -> None:
        slides = pl.segment("# T\n\nbody\n")
        catalog = [
            pl.LayoutInfo(
                name="Brand Colours", archetype="reference-non-content", is_content_layout=False
            ),
            pl.LayoutInfo(
                name="Bulleted Layout", archetype="bulleted-content", is_content_layout=True
            ),
        ]

        def fake_chat(prompt: str, _images, _schema):
            assert "Brand Colours" not in prompt  # never offered to the model
            return json.dumps({"choices": []})

        pl.select_layouts(slides, catalog, chat=fake_chat)
        assert slides[0].chosen_layout == "Bulleted Layout"

    def test_no_content_layouts_is_a_noop(self) -> None:
        slides = pl.segment("# T\n\nbody\n")
        catalog = [pl.LayoutInfo(name="Ref", is_content_layout=False)]
        pl.select_layouts(slides, catalog, chat=lambda *_: "unreachable")
        assert slides[0].chosen_layout is None


# ── VLM visual tie-break ─────────────────────────────────────────────────────
class TestVisualConfirm:
    def test_overrides_on_disagreement(self, tmp_path: Path) -> None:
        template = tmp_path / "t.pptx"
        template.write_bytes(b"")  # never opened; render is faked
        slides = pl.segment("# T\n\n- a\n- b\n")
        slides[0].chosen_layout = "Layout A"
        catalog = [
            pl.LayoutInfo(
                name="Layout A",
                archetype="bulleted-content",
                is_content_layout=True,
                example_pages=[1],
            ),
            pl.LayoutInfo(
                name="Layout B", archetype="photo-hero", is_content_layout=True, example_pages=[2]
            ),
        ]

        def fake_render(_template: Path, page: int) -> bytes | None:
            return f"png-{page}".encode()

        def fake_chat(_prompt: str, images, _schema):
            assert len(images) == 2
            return {"best": "Layout B", "why": "fits better"}

        pl.visual_confirm(slides, catalog, chat=fake_chat, render=fake_render, template=template)
        assert slides[0].chosen_layout == "Layout B"
        assert slides[0].chosen_why == "fits better"

    def test_alternates_come_from_the_slide_s_own_fit_ranking(self) -> None:
        """A tie-break is a choice between layouts that can hold the slide."""

        class _Report:
            def __init__(self, fits: bool) -> None:
                self.fits = fits

        catalog = {
            name: pl.LayoutInfo(
                name=name, archetype=archetype, is_content_layout=True, example_pages=[page]
            )
            for page, (name, archetype) in enumerate(
                [("Cover", "title-only"), ("Bullets", "bulleted-content"), ("Photo", "photo-hero")],
                start=1,
            )
        }
        slide = pl.segment("# T\n\n- a\n")[0]
        slide.chosen_layout = "Cover"
        # The fit model says the photo layout cannot hold this slide, so the
        # vision model must never be offered it.
        slide.fit_ranked = [
            (catalog["Cover"], _Report(True)),
            (catalog["Bullets"], _Report(True)),
            (catalog["Photo"], _Report(False)),
        ]
        names = [c.name for c in pl._alternates_for(slide, catalog, 3)]
        assert names == ["Cover", "Bullets"]
        # With no ranking at all, catalog order is the fallback it always was.
        slide.fit_ranked = []
        assert len(pl._alternates_for(slide, catalog, 3)) == 3

    def test_declines_an_override_that_would_flatten_the_deck(self, tmp_path: Path) -> None:
        """A tie-break is an arbitration, not a licence to redesign the deck."""
        template = tmp_path / "t.pptx"
        template.write_bytes(b"")
        slides = pl.segment("# One\n\n- a\n\n***\n\n## Two\n\n- b\n\n***\n\n## Three\n\n- c\n")
        # Two of the three slides already sit on Layout B, which is the share
        # past which a further override onto it stops being a tie-break.
        for slide, layout in zip(slides, ["Layout B", "Layout B", "Layout A"], strict=True):
            slide.chosen_layout = layout
        catalog = [
            pl.LayoutInfo(
                name="Layout A",
                archetype="bulleted-content",
                is_content_layout=True,
                example_pages=[1],
            ),
            pl.LayoutInfo(
                name="Layout B", archetype="photo-hero", is_content_layout=True, example_pages=[2]
            ),
        ]
        pl.visual_confirm(
            slides,
            catalog,
            chat=lambda *_: {"best": "Layout B", "why": "I like this one"},
            render=lambda _t, page: f"png-{page}".encode(),
            template=template,
        )
        assert slides[2].chosen_layout == "Layout A"  # declined: already 2 of 3
        # And the slides it was asked about that are already there stay there,
        # so declining never moves anything.
        assert [s.chosen_layout for s in slides[:2]] == ["Layout B", "Layout B"]

    def test_keeps_pick_when_render_unavailable(self, tmp_path: Path) -> None:
        template = tmp_path / "t.pptx"
        slides = pl.segment("# T\n\n- a\n")
        slides[0].chosen_layout = "Layout A"
        catalog = [
            pl.LayoutInfo(
                name="Layout A",
                archetype="bulleted-content",
                is_content_layout=True,
                example_pages=[1],
            ),
            pl.LayoutInfo(
                name="Layout B", archetype="photo-hero", is_content_layout=True, example_pages=[2]
            ),
        ]
        pl.visual_confirm(
            slides,
            catalog,
            chat=lambda *_: {"best": "Layout B"},
            render=lambda *_: None,
            template=template,
        )
        assert slides[0].chosen_layout == "Layout A"  # unchanged: nothing to compare

    def test_noop_without_template(self) -> None:
        slides = pl.segment("# T\n\n- a\n")
        slides[0].chosen_layout = "Layout A"
        pl.visual_confirm(slides, [], template=None)
        assert slides[0].chosen_layout == "Layout A"


class TestWhatALayoutActuallyIs:
    """A layout is described by the page that demonstrates it, when there is one."""

    def _template(self, tmp_path: Path) -> Path:
        """A template in the shape a Google Slides export has: no title, ever."""
        pptx = pytest.importorskip("pptx")
        from pptx.util import Inches, Pt

        prs = pptx.Presentation()
        layout = prs.slide_masters[0].slide_layouts[6]  # Blank
        page = prs.slides.add_slide(layout)
        heading = page.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8.0), Inches(1.2))
        heading.text_frame.text = "The heading of the slide"
        heading.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        body = page.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(8.0), Inches(3.0))
        body.text_frame.text = "Supporting copy the designer wrote"
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        out = tmp_path / "designer.pptx"
        prs.save(str(out))
        return out

    def test_the_demonstration_describes_the_layout_the_xml_does_not(self, tmp_path: Path) -> None:
        template = self._template(tmp_path)
        blank = next(c for c in pl.extract_layouts(template) if c.name == "Blank")
        kinds = [p.kind for p in blank.placeholders]
        # The Blank layout declares nothing at all; its example page shows a
        # heading over a body, and that is what a capacity model needs.
        assert "title" in kinds and "body" in kinds
        title = next(p for p in blank.placeholders if p.kind == "title")
        assert title.emu is not None and title.emu[1] < 1_000_000  # near the top

    def test_a_layout_nothing_demonstrates_keeps_its_own_declaration(self, tmp_path: Path) -> None:
        template = self._template(tmp_path)
        # Every other layout of the stock template has no example page, so it
        # is still described by its own XML — the fallback, unchanged.
        content = next(c for c in pl.extract_layouts(template) if c.name == "Title and Content")
        assert content.example_pages == []
        assert {p.kind for p in content.placeholders} >= {"ctrTitle"} or any(
            p.kind in ("title", "body") for p in content.placeholders
        )


class TestACoverIsMeasuredAsOne:
    def test_segment_marks_the_cover_and_the_closing(self) -> None:
        slides = pl.segment("# Cover\n\n***\n\n## Middle\n\n- a\n\n***\n\n## Thanks\n")
        assert [s.title_only for s in slides] == [True, False, True]

    def test_a_single_slide_deck_is_both(self) -> None:
        assert pl.segment("# Only\n\n- a\n")[0].title_only is True
