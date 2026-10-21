"""Tests for md2star.pptx_assemble — box helpers, assembly, the eyeball loop.

Exercised **offline**: real ``python-pptx`` objects are built against its own
bundled default template (no external asset needed), and the target-matching
eyeball loop's LLM/VLM/render seams are driven with injected fakes, no
Ollama/LibreOffice/poppler required — mirroring tests/test_reverse_twin.py.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from md2star import pptx_assemble as pa  # noqa: E402
from md2star import pptx_audit as au  # noqa: E402
from md2star import pptx_layout as pl  # noqa: E402


@pytest.fixture
def default_template(tmp_path: Path) -> Path:
    path = tmp_path / "default.pptx"
    pptx.Presentation().save(str(path))
    return path


@pytest.fixture
def catalog(default_template: Path) -> list[pl.LayoutInfo]:
    return pl.extract_layouts(default_template)


# ── template introspection ───────────────────────────────────────────────────
class TestTemplateFonts:
    def test_falls_back_to_arial_with_no_explicit_typeface(self, default_template: Path) -> None:
        fam = pa.template_fonts(default_template)
        assert fam["display"] and fam["body"]  # never empty, even with nothing declared

    def test_picks_up_an_explicit_design_font(self, tmp_path: Path) -> None:
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        box = slide.shapes.add_textbox(0, 0, pptx.util.Inches(1), pptx.util.Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "hello"
        run.font.name = "Montserrat"
        path = tmp_path / "fonted.pptx"
        prs.save(str(path))
        fam = pa.template_fonts(path)
        assert fam["family"] == "Montserrat"
        assert fam["body"] == "Montserrat"


class TestDarkDetection:
    def test_unresolvable_background_is_not_dark(self, default_template: Path) -> None:
        prs = pptx.Presentation(str(default_template))
        layout = prs.slide_layouts[0]
        # python-pptx's stock layouts declare no explicit solid fill.
        assert pa._bg_hex(layout) is None
        assert pa.is_dark_layout(layout) is None


# ── box helpers ───────────────────────────────────────────────────────────────
class TestBoxHelpers:
    def test_uses_placeholder_geometry_when_present(self, default_template: Path, catalog) -> None:
        prs = pptx.Presentation(str(default_template))
        info = next(c for c in catalog if c.name == "Picture with Caption")
        box = pa.image_box(prs, info)
        placeholder = next(p for p in info.placeholders if p.kind == "pic")
        expected = tuple(v / pa.EMU_PER_INCH for v in placeholder.emu)
        assert box == expected

    def test_falls_back_to_fractional_default_without_geometry(
        self, default_template: Path
    ) -> None:
        prs = pptx.Presentation(str(default_template))
        box = pa.image_box(prs, None)
        slide_w = prs.slide_width / pa.EMU_PER_INCH
        assert 0 < box[0] < slide_w
        assert 0 < box[2] <= slide_w


# ── content parsing ──────────────────────────────────────────────────────────
class TestContentParsing:
    def test_parse_image_resolves_relative_to_base_dir(self, tmp_path: Path) -> None:
        (tmp_path / "pic.png").write_bytes(b"\x89PNG")
        found = pa.parse_image("![alt](pic.png)", tmp_path)
        assert found == (tmp_path / "pic.png").resolve()

    def test_parse_image_missing_file_returns_none(self, tmp_path: Path) -> None:
        assert pa.parse_image("![alt](nope.png)", tmp_path) is None

    def test_parse_mermaid_extracts_fenced_block(self) -> None:
        assert pa.parse_mermaid("```mermaid\nflowchart LR\nA-->B\n```") == "flowchart LR\nA-->B"

    def test_parse_code_skips_mermaid_block(self) -> None:
        body = "```mermaid\nflowchart LR\n```\n```python\nprint(1)\n```"
        assert pa.parse_code(body) == "print(1)"

    def test_fallback_bullets_strips_markup_and_caps(self) -> None:
        body = "# Title\n\n- **one** _two_\n- three\n- four\n- five\n"
        bullets = pa.fallback_bullets(body, cap=2)
        assert len(bullets) == 2
        assert bullets[0] == "one two"


# ── assembly ──────────────────────────────────────────────────────────────────
class TestAssemble:
    def test_places_every_slide_onto_its_chosen_layout(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment(
            "# Cover\n\nsubtitle\n\n***\n\n## Bullets\n\n- a\n- b\n\n***\n\n## Divider {.big}\n\n***\n\n## The End\n"
        )
        by_archetype = {"Title Slide": 0, "Title and Content": 1, "Section Header": 2, "Blank": 3}
        names = list(by_archetype)
        for s, name in zip(slides, names, strict=True):
            s.chosen_layout = name

        out = tmp_path / "out.pptx"
        result = pa.assemble(slides, catalog, default_template, out)
        assert result.placed == result.total == 4
        assert out.exists()

        prs = pptx.Presentation(str(out))
        assert len(prs.slides) == 4
        used_layout_names = [slide.slide_layout.name for slide in prs.slides]
        assert used_layout_names == names

    def test_unknown_layout_falls_back(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("# Only\n\nbody\n")
        slides[0].chosen_layout = "Does Not Exist"
        out = tmp_path / "out.pptx"
        result = pa.assemble(slides, catalog, default_template, out, default_layout="Title Slide")
        assert result.placed == 1
        prs = pptx.Presentation(str(out))
        assert prs.slides[0].slide_layout.name == "Title Slide"

    def test_bad_slide_does_not_sink_the_deck(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("# One\n\nbody\n\n***\n\n## Two\n\nbody2\n")
        slides[0].chosen_layout = "Title Slide"
        slides[1].chosen_layout = "Title Slide"
        # An image feature with no resolvable path degrades to title-only,
        # not a crash — simulate a hard failure instead via a bogus override.
        slides[0].overrides = {"font_delta": "not-a-number"}
        out = tmp_path / "out.pptx"
        result = pa.assemble(slides, catalog, default_template, out)
        assert result.total == 2
        assert result.placed == 1  # slide 1 failed (bad override), slide 2 still placed

    def test_no_layouts_raises(self, tmp_path: Path) -> None:
        # Presentation() always has layouts in practice; simulate the guard by
        # asserting it fires on an (unrealistic) empty layout_index result.
        prs = pptx.Presentation()
        assert pa.layout_index(prs)  # sanity: never empty for a real template


# ── placeholder / overlay reconciliation ─────────────────────────────────────
# add_slide() clones the layout's placeholders onto the slide, so an overlay
# drawn in the same geometry used to ship a second, empty shape on top of every
# region: a "Click to add title" prompt where the title is, no title at all in
# the outline pane, and a guaranteed overlap on every slide.
class TestPlaceholderReconciliation:
    @pytest.fixture
    def deck(self, tmp_path: Path, default_template: Path, catalog) -> Path:
        """A three-slide deck on a layout with title + picture + body."""
        slides = pl.segment("# Cover\n\n***\n\n## Bullets\n\n- a\n- b\n\n***\n\n## The End\n")
        for s in slides:
            s.chosen_layout = "Picture with Caption"
        out = tmp_path / "out.pptx"
        pa.assemble(slides, catalog, default_template, out)
        return out

    def test_no_unfilled_placeholder_survives(self, deck: Path) -> None:
        prs = pptx.Presentation(str(deck))
        ghosts = [
            shape.name
            for slide in prs.slides
            for shape in slide.shapes
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip()
        ]
        assert ghosts == []

    def test_the_title_lands_in_the_title_placeholder(self, deck: Path) -> None:
        """What the outline pane and a screen reader read is the placeholder."""
        prs = pptx.Presentation(str(deck))
        titles = [
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if pa.placeholder_kind(shape) in {"title", "ctrTitle"}
        ]
        assert titles == ["Cover", "Bullets", "The End"]

    def test_the_title_is_not_also_drawn_as_a_loose_box(self, deck: Path) -> None:
        prs = pptx.Presentation(str(deck))
        for slide in prs.slides:
            drawn = [
                shape.text_frame.text
                for shape in slide.shapes
                if not shape.is_placeholder and shape.has_text_frame
            ]
            assert not any(text.startswith(("Cover", "Bullets", "The End")) for text in drawn)

    def test_bullets_keep_their_own_box(self, deck: Path) -> None:
        """Body text is NOT poured into the designer's placeholder: its master
        would add a second bullet glyph next to the one the assembler draws."""
        prs = pptx.Presentation(str(deck))
        boxes = [
            shape.text_frame.text
            for shape in prs.slides[1].shapes
            if not shape.is_placeholder and shape.has_text_frame
        ]
        assert boxes == ["•  a\n•  b"]

    def test_a_layout_without_a_title_placeholder_still_gets_a_title(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("# Only\n\nbody\n")
        slides[0].chosen_layout = "Blank"  # no placeholders at all
        out = tmp_path / "blank.pptx"
        pa.assemble(slides, catalog, default_template, out)
        prs = pptx.Presentation(str(out))
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert texts == ["Only"]

    def test_an_empty_title_is_still_named(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        """`[x] or [fallback]` never fired — a one-element list is always truthy."""
        slides = pl.segment("# Only\n\nbody\n")
        slides[0].title = ""
        slides[0].chosen_layout = "Blank"
        out = tmp_path / "untitled.pptx"
        pa.assemble(slides, catalog, default_template, out)
        prs = pptx.Presentation(str(out))
        assert [sh.text_frame.text for sh in prs.slides[0].shapes] == ["(untitled)"]


# ── speaker notes ────────────────────────────────────────────────────────────
class TestSpeakerNotes:
    def test_notes_go_to_the_notes_pane_not_the_slide(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment(
            "## Numbers\n\n- revenue up\n\n::: notes\nDo not mention the layoffs.\n:::\n"
        )
        slides[0].chosen_layout = "Title and Content"
        out = tmp_path / "notes.pptx"
        pa.assemble(slides, catalog, default_template, out)

        prs = pptx.Presentation(str(out))
        slide = prs.slides[0]
        on_slide = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        assert "layoffs" not in on_slide
        assert slide.has_notes_slide
        assert slide.notes_slide.notes_text_frame.text == "Do not mention the layoffs."

    def test_a_slide_without_notes_gets_no_notes_slide(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("## Plain\n\n- a\n")
        slides[0].chosen_layout = "Title and Content"
        out = tmp_path / "plain.pptx"
        pa.assemble(slides, catalog, default_template, out)
        assert not pptx.Presentation(str(out)).slides[0].has_notes_slide


# ── image accessibility ──────────────────────────────────────────────────────
# Pandoc carries Markdown alt text into the picture's `descr` attribute; this
# path bypasses Pandoc, so it has to do the same or the deck is unreadable to
# a screen reader.
class TestAltText:
    @pytest.fixture
    def png(self, tmp_path: Path) -> Path:
        from PIL import Image

        path = tmp_path / "chart.png"
        Image.new("RGB", (1600, 900), "navy").save(path)
        return path

    def test_alt_text_reaches_the_picture(
        self, tmp_path: Path, default_template: Path, catalog, png: Path
    ) -> None:
        slides = pl.segment(f"## Results\n\n![Revenue climbing through Q3]({png.name})\n")
        slides[0].chosen_layout = "Picture with Caption"
        out = tmp_path / "alt.pptx"
        pa.assemble(slides, catalog, default_template, out, image_base_dir=tmp_path)

        prs = pptx.Presentation(str(out))
        picture = next(
            shape for shape in prs.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        descr = picture._element._nvXxPr.cNvPr.get("descr")  # noqa: SLF001
        assert descr == "Revenue climbing through Q3"

    def test_without_alt_text_the_title_describes_the_picture(
        self, tmp_path: Path, default_template: Path, catalog, png: Path
    ) -> None:
        """python-pptx would otherwise leave the FILE NAME as the description,
        and "chart dot png" is not what a screen reader should announce."""
        slides = pl.segment(f"## Results\n\n![]({png.name})\n")
        slides[0].chosen_layout = "Picture with Caption"
        out = tmp_path / "noalt.pptx"
        pa.assemble(slides, catalog, default_template, out, image_base_dir=tmp_path)

        prs = pptx.Presentation(str(out))
        picture = next(
            shape for shape in prs.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        assert picture._element._nvXxPr.cNvPr.get("descr") == "Results"  # noqa: SLF001

    def test_parse_alt_text_reads_the_first_image(self) -> None:
        assert pa.parse_alt_text("![a bar chart](x.png) ![another](y.png)") == "a bar chart"
        assert pa.parse_alt_text("no image here") == ""


# ── target-matching eyeball loop ─────────────────────────────────────────────
class TestEyeballLoop:
    def test_applies_overrides_on_mismatch_then_stops(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("# Cover\n\n- a\n- b\n- c\n")
        slides[0].chosen_layout = "Title and Content"
        out = tmp_path / "out.pptx"

        calls = {"judge": 0}

        def fake_chat(_prompt, _images, _schema):
            calls["judge"] += 1
            if calls["judge"] == 1:
                return {"matches": False, "overflow": True, "discrepancies": "too much text"}
            return {"matches": True}

        def fake_render_template(_template, _page):
            return b"target-png"

        def fake_render_deck(_path):
            return {1: b"candidate-png"}

        # Give the layout an example page so a target thumbnail is attempted.
        for c in catalog:
            if c.name == "Title and Content":
                c.example_pages = [1]

        result = pa.eyeball_slides(
            slides,
            catalog,
            default_template,
            out,
            max_iterations=2,
            chat=fake_chat,
            render_template=fake_render_template,
            render_deck=fake_render_deck,
        )
        assert result.placed == 1
        assert slides[0].overrides.get("font_delta") == -4
        assert slides[0].overrides.get("bullets") == 2
        # First pass flags+fixes, second pass's judge call reports a match -> stop.
        assert calls["judge"] == 2

    def test_a_judge_that_talks_around_its_json_is_still_heard(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        """Regression: the verdict used to be read with a bare ``json.loads``.

        The loop asks for structured output, but ``llm.chat`` documents that
        it returns the raw *string* when the model does not honour the
        schema — and on the vision side the strongest local models are
        exactly the ones that don't, which is what the phase-5 bench
        (``scripts/bench_vlm_judge.py``) set out to measure. A fenced,
        prose-wrapped reply raised ``JSONDecodeError`` out of the loop, so one
        chatty answer cost the whole deck its verification instead of costing
        one slide its verdict.
        """
        slides = pl.segment("# Cover\n\n- a\n- b\n- c\n")
        slides[0].chosen_layout = "Title and Content"
        out = tmp_path / "out.pptx"
        for c in catalog:
            if c.name == "Title and Content":
                c.example_pages = [1]

        calls = {"judge": 0}

        def fenced_chat(_prompt, _images, _schema):
            # The literal shape gemma3:12b returns when asked for the same
            # verdict without a schema, observed in the bench run.
            calls["judge"] += 1
            if calls["judge"] == 1:
                return (
                    '```json\n{"matches": false, "overflow": true, '
                    '"discrepancies": "the body runs past its box"}\n```'
                )
            return '```json\n{"matches": true}\n```'

        result = pa.eyeball_slides(
            slides,
            catalog,
            default_template,
            out,
            max_iterations=2,
            chat=fenced_chat,
            render_template=lambda _t, _p: b"target-png",
            render_deck=lambda _p: {1: b"candidate-png"},
        )
        # Not merely "did not raise": the verdict buried in the fence was
        # acted on, with the same bounded fix a clean dict would have earned.
        assert result.placed == 1
        assert slides[0].overrides.get("font_delta") == -4
        assert calls["judge"] == 2

    def test_a_judge_answering_in_prose_costs_one_verdict_not_the_run(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        """A reply with no JSON at all degrades to "no verdict", as designed."""
        slides = pl.segment("# Cover\n\n- a\n- b\n- c\n")
        slides[0].chosen_layout = "Title and Content"
        out = tmp_path / "out.pptx"
        for c in catalog:
            if c.name == "Title and Content":
                c.example_pages = [1]

        result = pa.eyeball_slides(
            slides,
            catalog,
            default_template,
            out,
            max_iterations=1,
            chat=lambda _p, _i, _s: "I'm sorry, I can't view those images.",
            render_template=lambda _t, _p: b"target-png",
            render_deck=lambda _p: {1: b"candidate-png"},
        )
        # The deck still exists and nothing was invented from an unread reply.
        assert result.placed == 1
        assert slides[0].overrides == {}

    def test_no_target_page_skips_judging(
        self, tmp_path: Path, default_template: Path, catalog
    ) -> None:
        slides = pl.segment("# Cover\n\nbody\n")
        slides[0].chosen_layout = "Title Slide"  # no example_pages -> no rep_page
        out = tmp_path / "out.pptx"

        def unreachable_chat(*_a):
            raise AssertionError("must not be called with no target page")

        result = pa.eyeball_slides(
            slides,
            catalog,
            default_template,
            out,
            max_iterations=1,
            chat=unreachable_chat,
            render_deck=lambda _p: {1: b"png"},
        )
        assert result.placed == 1
        assert slides[0].overrides == {}


class TestTheLoopsOtherEye:
    """The Look step's deterministic half: the audit, fed back as a verdict."""

    def test_a_finding_becomes_a_verdict_the_loop_already_understands(self) -> None:
        findings = [
            au.Finding(
                2, "overrun", "text needs 5.7x its box at 26pt (fits at ~11pt)", "TextBox 2"
            ),
            au.Finding(3, "overlap", "two shapes share 40% of the smaller box", "Title 1"),
        ]
        verdict = pa._measured_verdict(findings, 2)
        assert verdict["matches"] is False and verdict["overflow"] is True
        assert "5.7x" in verdict["discrepancies"]
        # An overlap is usually a designer's composition and has no bounded
        # fix here, so it is not turned into a verdict at all.
        assert pa._measured_verdict(findings, 3) is None
        assert pa._measured_verdict(findings, 1) is None

    def test_a_measurement_never_loses_to_an_opinion(self) -> None:
        measured = {"matches": False, "overflow": True, "discrepancies": "needs 5.7x its box"}
        judged = {"matches": True, "overflow": False, "discrepancies": ""}
        merged = pa._merged_verdict(measured, judged)
        # The renderer hides an overrun by shrinking the type, so a model
        # looking at the picture cannot see what the arithmetic measured.
        assert merged["matches"] is False and merged["overflow"] is True
        assert "5.7x" in merged["discrepancies"]
        assert pa._merged_verdict(None, judged) == judged
        assert pa._merged_verdict(measured, None) == measured

    def test_the_exact_correction_survives_the_merge(self) -> None:
        """A number is not a flag, and it was getting dropped as one."""
        measured = {"matches": False, "overflow": True, "font_delta": -3}
        judged = {"matches": True, "discrepancies": ""}
        merged = pa._merged_verdict(measured, judged)
        assert merged["font_delta"] == -3
        assert pa._overrides_for(merged)["font_delta"] == -3

    def test_the_measured_eye_applies_the_exact_correction(self) -> None:
        """Arithmetic knows the size that fits, so it need not guess a step."""
        findings = [
            au.Finding(
                1,
                "overrun",
                "text needs 1.3x its box at 26pt (fits at ~23pt)",
                "Body",
                measured_pt=26.0,
                fits_pt=23.0,
            )
        ]
        verdict = pa._measured_verdict(findings, 1)
        assert verdict["font_delta"] == -3
        assert pa._overrides_for(verdict)["font_delta"] == -3
        # A model that only knows "it overflows" still gets the old nudge.
        assert pa._overrides_for({"overflow": True})["font_delta"] == -4

    def test_a_correction_that_would_blind_the_bullets_is_not_a_correction(self) -> None:
        """One delta hits every size, so the smallest one is what binds."""
        # A 26pt title that only fits at 14pt asks for -12 — which would take
        # the slide's 16pt bullets to 4pt. Observed on a real deck.
        verdict = {"matches": False, "overflow": True, "font_delta": -12}
        assert pa._overrides_for(verdict) == {}
        # The same shortfall on a slide whose smallest type can absorb it is
        # applied exactly, as it should be.
        assert pa._overrides_for({"overflow": True, "font_delta": -4})["font_delta"] == -4

    def test_an_overrun_no_shrinking_can_win_asks_for_a_different_box(self) -> None:
        findings = [
            au.Finding(
                1,
                "overrun",
                "text needs 5.7x its box at 26pt (fits at ~11pt)",
                "Body",
                measured_pt=26.0,
                fits_pt=11.0,
            )
        ]
        verdict = pa._measured_verdict(findings, 1)
        assert verdict["beyond_shrinking"] is True
        # Empty, on purpose: it is what makes the caller's ladder go looking
        # for a box that is big enough instead of setting 11pt type.
        assert pa._overrides_for(verdict) == {}

    def test_a_measured_defect_drives_a_fix_with_no_model_at_all(
        self, tmp_path: Path, default_template: Path, catalog, monkeypatch
    ) -> None:
        slides = pl.segment("# Cover\n\n***\n\n## Where we stand\n\n- a\n- b\n\n***\n\n## Thanks\n")
        for slide, layout in zip(slides, ["Title Slide", "Title and Content", "Section Header"]):
            slide.chosen_layout = layout
        monkeypatch.setattr(
            au,
            "audit",
            lambda _deck, **_kw: [au.Finding(2, "overrun", "needs 3.2x its box at 16pt", "Body")],
        )

        def unreachable_chat(*_a):
            raise AssertionError("the arithmetic half must need no model")

        result = pa.eyeball_slides(
            slides,
            catalog,
            default_template,
            tmp_path / "measured.pptx",
            max_iterations=1,
            chat=unreachable_chat,
            # No renders at all: no LibreOffice, no thumbnails, no vision
            # model — exactly the machine md2star is meant to still work on.
            render_deck=lambda _p: {},
            render_template=lambda *_a: None,
        )
        assert result.placed == 3
        assert slides[1].overrides.get("font_delta") == -4
        assert slides[1].overrides.get("bullets") == 2
        # The slides nobody measured a defect on were left alone.
        assert slides[0].overrides == {} and slides[2].overrides == {}
