#!/usr/bin/env python3
"""Bench the eyeball loop's vision judge: strict-JSON vs. lenient parsing.

Module summary
--------------
Phase 5 of ``docs/design/pptx-template-intelligence.md`` asks a question the
design could not settle on paper: the slide eyeball loop
(:func:`md2star.pptx_assemble.eyeball_slides`) needs a ``{matches, overflow,
low_contrast, too_empty, image_too_small, discrepancies}`` verdict from a
vision model, and there are two ways to get one.

*Strict* asks Ollama for structured output against :data:`_JUDGE_SCHEMA`.
``gemma3`` honours it; ``qwen3-vl`` — which scores higher on vision — does not
reliably, so the design flagged a **lenient** second mode: ask for the same
JSON in prose and recover it with a tolerant parse, the way
:func:`md2star.reverse_diagrams._parse_verdict` already recovers the diagram
loop's verdicts. Which to ship is an empirical question about two things at
once: *can the reply be consumed at all*, and *is the judgement any good*.

The corpus answers the second honestly because its ground truth is **measured,
not opined**. Each case is one slide the project's own assembler built on a
designer layout, then deliberately broken in exactly one way; the arithmetic
audit (:func:`md2star.pptx_audit.audit`) is then asked to confirm the label
before the case is allowed into the bench. A case the audit cannot confirm is
reported as unconfirmed rather than quietly trusted.

One case is special and the bench says so out loud. ``hidden_overrun`` sets
text far too large for its box: the audit sees the overrun in the XML, but
PowerPoint and LibreOffice *auto-shrink* the text, so the render looks clean.
No vision model can see it, and scoring one as wrong for missing it would
measure the renderer rather than the model. It is carried as an unscored
control, and it is the empirical reason
:func:`md2star.pptx_assemble._merged_verdict` lets arithmetic win on overflow.

Usage example
-------------
>>> # Full bench, both models, both modes (needs Ollama + LibreOffice):
>>> #   python scripts/bench_vlm_judge.py
>>> # One arm, keeping the renders to look at:
>>> #   python scripts/bench_vlm_judge.py --models qwen3-vl:4b --keep /tmp/bench
>>> # Machine-readable, for a follow-up round to diff against:
>>> #   python scripts/bench_vlm_judge.py --json results.json

Author
------
[Warith HARCHAOUI](https://linkedin.com/in/warith-harchaoui/)
"""

from __future__ import annotations

import argparse
import json
import re
import statistics
import sys
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path

# The bench drives md2star's real stages, never a reimplementation of them:
# what it measures has to be what the product would actually get back.
from md2star.pptx_assemble import _JUDGE_PROMPT, _JUDGE_SCHEMA, assemble
from md2star.pptx_audit import audit
from md2star.pptx_layout import (
    Slide,
    build_catalog,
    default_chat,
    default_render,
    render_deck_pages,
)

#: The template the design's phase plan benches against, and the layout within
#: it. ``Segmentation 2A`` earns the job on two counts: a plain three-bullet
#: slide assembles onto it with *no* audit finding (so the clean control is
#: genuinely clean), and the designer demonstrates it on a real page, so the
#: target image is their own composition rather than a bare layout.
DEFAULT_TEMPLATE = Path("assets/templates/pptx/Jellysmack-Presentation-Template.pptx")
BENCH_LAYOUT = "Segmentation 2A"
BENCH_PAGE = 36

#: Content held fixed across every case, so the only thing that varies between
#: two candidate renders is the defect injected into one of them.
BENCH_TITLE = "Quarterly operating review"
BENCH_BODY = "- Revenue up 12 percent\n- Churn flat for two quarters\n- Two hires still pending"


@dataclass(frozen=True)
class Truth:
    """What a case is, and how we know it.

    Attributes
    ----------
    matches : bool
        Whether a correct judge should call the candidate clean.
    flags : frozenset of str
        The verdict flags a correct judge should raise, in the schema's own
        vocabulary (``overflow`` / ``low_contrast`` / ``too_empty`` /
        ``image_too_small``).
    audit_kind : str or None
        The :mod:`md2star.pptx_audit` finding that independently confirms the
        label, or ``None`` for a defect arithmetic is blind to — which is
        precisely the kind a vision model is *for*.
    visible : bool
        ``False`` when the renderer hides the defect (see ``hidden_overrun``).
        Unscored: the bench reports it, it never counts it against a model.
    """

    matches: bool
    flags: frozenset[str] = frozenset()
    audit_kind: str | None = None
    visible: bool = True


#: The corpus, in the order it is built onto the deck. Keep ``clean`` first:
#: a judge that flags it is crying wolf, which costs the loop a needless
#: layout swap, and that false-alarm rate is reported separately from accuracy.
CORPUS: dict[str, Truth] = {
    "clean": Truth(matches=True, audit_kind=None),
    # Text three times too big for its box. The audit reads the XML and sees
    # it; both renderers auto-shrink it away. Carried, never scored.
    "hidden_overrun": Truth(
        matches=False, flags=frozenset({"overflow"}), audit_kind="overrun", visible=False
    ),
    # The same overrun with autofit and wrapping switched off, so it actually
    # runs out of its box and off the slide — the visible half of "overflow".
    "spill": Truth(matches=False, flags=frozenset({"overflow"}), audit_kind="overrun"),
    "off_slide": Truth(matches=False, flags=frozenset({"overflow"}), audit_kind="off-slide"),
    "low_contrast": Truth(
        matches=False, flags=frozenset({"low_contrast"}), audit_kind="low-contrast"
    ),
    # No arithmetic check exists for "the composition collapsed" — this is the
    # case that only an eye can settle, and the reason the loop has one.
    "too_empty": Truth(matches=False, flags=frozenset({"too_empty"}), audit_kind=None),
}


@dataclass
class Case:
    """One built, rendered and label-confirmed bench case."""

    name: str
    truth: Truth
    candidate: bytes
    #: ``True`` when the audit reported the finding the label promised. A
    #: case with ``audit_kind=None`` is confirmed by construction instead.
    confirmed: bool


@dataclass
class Reply:
    """One judge call, scored along both consumption paths."""

    case: str
    raw: object
    seconds: float
    #: What the *current* consumer in ``eyeball_slides`` gets. ``None`` means
    #: it got nothing; ``"raise"`` in :attr:`crash` means it got an exception.
    strict: dict | None = None
    crash: str = ""
    #: What a tolerant parse recovers from the same reply.
    lenient: dict | None = None


# ── corpus ───────────────────────────────────────────────────────────────────
def _inject(name: str, shape, slide_width: int) -> None:
    """Break *shape* in exactly the one way *name* calls for.

    Each injection is the smallest edit that produces the defect: nothing else
    about the slide changes, so a judge's verdict is about that defect alone.
    """
    # Imported here rather than at module scope: python-pptx is an opt-in
    # extra, and `--help` should work on a machine that never installed it.
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.enum.text import MSO_AUTO_SIZE  # noqa: PLC0415
    from pptx.util import Emu, Pt  # noqa: PLC0415

    frame = shape.text_frame
    runs = [run for para in frame.paragraphs for run in para.runs]
    if name == "hidden_overrun":
        for run in runs:
            run.font.size = Pt(54)
    elif name == "spill":
        # Autofit is what hides an overrun; switching it (and wrapping) off is
        # what turns the same edit into something an eye can actually see.
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = False
        for run in runs:
            run.font.size = Pt(44)
            run.text = run.text + " and the same three constraints restated again"
    elif name == "off_slide":
        shape.left = Emu(int(slide_width * 0.88))
    elif name == "low_contrast":
        # Near-black on this layout's black field. The value matters: an
        # earlier near-*white* injection was perfectly legible here, which is
        # the whole reason the audit resolves the backdrop instead of assuming.
        for run in runs:
            run.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    elif name == "too_empty":
        frame.text = "•  Yes"


def build_corpus(template: Path, workdir: Path) -> tuple[list[Case], bytes]:
    """Assemble, break, render and label-check every case.

    Returns
    -------
    tuple
        The confirmed cases, and the target PNG — the designer's own example
        page, which is what the loop shows a judge as "what this should
        look like".
    """
    from pptx import Presentation  # noqa: PLC0415 — opt-in extra, see _inject

    catalog = build_catalog(template, use_vision=False)
    names = list(CORPUS)
    # Slide 1 and slide n are forced title-only by the assembler (a cover and a
    # closing carry a title and nothing else), so the cases are padded to sit
    # strictly inside that range and all get real content.
    slides = [Slide(index=1, title="Cover", title_only=True, chosen_layout=BENCH_LAYOUT)]
    slides += [
        Slide(
            index=i + 2,
            title=BENCH_TITLE,
            body=BENCH_BODY,
            features=["bullets"],
            chosen_layout=BENCH_LAYOUT,
        )
        for i in range(len(names))
    ]
    slides.append(
        Slide(index=len(names) + 2, title="Thank you", title_only=True, chosen_layout=BENCH_LAYOUT)
    )

    deck = workdir / "bench.pptx"
    assemble(slides, catalog, template, deck)

    prs = Presentation(str(deck))
    for name, slide in zip(names, list(prs.slides)[1:], strict=False):
        # The assembler's own body overlay is the one carrying bullet glyphs;
        # that is what identifies it without depending on shape names.
        body = next(sh for sh in slide.shapes if sh.has_text_frame and "•" in sh.text_frame.text)
        _inject(name, body, prs.slide_width)
    prs.save(str(deck))

    # The label check: ask arithmetic whether it sees what we claim we broke.
    seen: dict[int, set[str]] = {}
    for finding in audit(deck):
        seen.setdefault(finding.slide, set()).add(finding.kind)
    pages = render_deck_pages(deck)

    cases: list[Case] = []
    for i, name in enumerate(names, start=2):
        truth = CORPUS[name]
        kinds = seen.get(i, set())
        if truth.audit_kind is None:
            # Nothing for arithmetic to confirm — but a *clean* control must
            # still come back clean, or the control is not one.
            confirmed = not kinds if truth.matches else True
        else:
            confirmed = truth.audit_kind in kinds
        cases.append(Case(name=name, truth=truth, candidate=pages[i], confirmed=confirmed))
    return cases, default_render()(template, BENCH_PAGE)


# ── the two consumption paths ────────────────────────────────────────────────
def consume_strict(raw: object) -> tuple[dict | None, str]:
    """Parse *raw* exactly the way :func:`eyeball_slides` does today.

    This is deliberately not a tolerant parse: the bench's first question is
    what the shipping consumer gets, so it runs the shipping consumer's code
    shape — ``dict`` straight through, ``str`` through :func:`json.loads`,
    anything else dropped. The returned string is non-empty when that path
    *raised*, which is itself a result.
    """
    if isinstance(raw, dict):
        return raw, ""
    if isinstance(raw, str):
        try:
            parsed = json.loads(raw)
        except (json.JSONDecodeError, TypeError) as exc:
            return None, f"{type(exc).__name__}: {exc}"
        return (parsed if isinstance(parsed, dict) else None), ""
    return None, ""


def consume_lenient(raw: object) -> dict | None:
    """Recover a verdict from a reply that merely *contains* JSON.

    The same stance :func:`md2star.reverse_diagrams._parse_verdict` already
    takes for the diagram loop: models wrap their JSON in prose or fences, and
    the first ``{...}`` span is almost always the verdict. Returns ``None``
    only when there is no object in there at all.
    """
    if isinstance(raw, dict):
        return raw
    if not isinstance(raw, str):
        return None
    # Greedy to the last brace: the verdict object is the outermost one, and a
    # nested object would truncate a non-greedy match halfway through.
    span = re.search(r"\{.*\}", raw, re.DOTALL)
    if not span:
        return None
    try:
        parsed = json.loads(span.group(0))
    except (json.JSONDecodeError, TypeError):
        return None
    return parsed if isinstance(parsed, dict) else None


# ── running one arm ──────────────────────────────────────────────────────────
def run_arm(
    model: str, mode: str, cases: list[Case], target: bytes, *, prompt: str = _JUDGE_PROMPT
) -> list[Reply]:
    """Judge every case with one (model, mode) pair.

    *mode* is ``"strict"`` (ask Ollama for structured output against the
    schema) or ``"lenient"`` (ask for the same JSON in prose, recover it
    afterwards) — the two halves of the question phase 5 poses.
    """
    chat = default_chat(kind="vlm", model=model)
    # The lenient arm still has to *ask* for the JSON; what it drops is the
    # schema constraint on the transport, not the instruction in the prompt.
    schema = _JUDGE_SCHEMA if mode == "strict" else None
    replies: list[Reply] = []
    for case in cases:
        # Target first, candidate second — the order the prompt describes and
        # the order the loop uses. Swapping them silently inverts every
        # verdict, so the bench must not "tidy" it.
        started = time.monotonic()
        raw = chat(prompt, [target, case.candidate], schema)
        elapsed = time.monotonic() - started
        # Both consumption paths are run on the *same* reply, never on two
        # separate calls: a model asked twice answers twice, and the question
        # here is what two parsers make of one answer.
        strict, crash = consume_strict(raw)
        replies.append(
            Reply(
                case=case.name,
                raw=raw,
                seconds=elapsed,
                strict=strict,
                crash=crash,
                lenient=consume_lenient(raw),
            )
        )
        print(f"    {case.name:16} {elapsed:6.1f}s  {_oneline(raw)}", flush=True)
    return replies


def _oneline(raw: object, width: int = 88) -> str:
    """Collapse a reply to one readable line for the live log."""
    text = json.dumps(raw) if isinstance(raw, dict) else str(raw)
    return " ".join(text.split())[:width]


# ── scoring ──────────────────────────────────────────────────────────────────
@dataclass
class Score:
    """What one arm achieved, over the scored (visible) cases only."""

    model: str
    mode: str
    # Counted apart from every judgement metric below, because it is not one:
    # Ollama 500ing or timing out says nothing about whether the model can
    # see a clipped slide, and folding the two together would read as a
    # model being wrong when it was never asked.
    transport_failures: int = 0  # the call came back with nothing at all
    usable_strict: int = 0  # replies today's consumer turns into a verdict
    crashes: int = 0  # replies today's consumer *raises* on
    usable_lenient: int = 0  # replies a tolerant parse recovers
    correct_matches: int = 0  # right call on "is this slide clean?"
    flag_hits: int = 0  # ...and named the right defect
    false_alarms: int = 0  # flagged the clean control
    scored: int = 0
    latencies: list[float] = field(default_factory=list)

    @property
    def answered(self) -> int:
        """Scored cases the model actually replied to — the honest denominator."""
        return self.scored - self.transport_failures

    @property
    def median_s(self) -> float:
        return statistics.median(self.latencies) if self.latencies else 0.0


def score_arm(model: str, mode: str, cases: list[Case], replies: list[Reply]) -> Score:
    """Turn one arm's replies into counts, ignoring the unscorable case."""
    by_name = {c.name: c for c in cases}
    out = Score(model=model, mode=mode)
    for reply in replies:
        case = by_name[reply.case]
        # The invisible case is carried through the run and printed, but it
        # cannot discriminate between models, so it is kept out of the counts.
        if not case.truth.visible:
            continue
        out.scored += 1
        out.latencies.append(reply.seconds)
        # `default_chat` swallows every transport error into ``None``, which
        # is the one reply that is not an answer. Separate it before anything
        # else, or a flaky daemon reads as a blind model.
        if reply.raw is None:
            out.transport_failures += 1
            continue
        # A crash and an unusable reply are counted apart on purpose: both
        # cost the verdict, but only one of them takes the stage down with it.
        if reply.crash:
            out.crashes += 1
        if reply.strict is not None:
            out.usable_strict += 1
        # Judgement quality is scored off the *lenient* parse throughout, so
        # the two arms are compared on what their model saw rather than on
        # which of them happened to be readable by today's consumer.
        verdict = reply.lenient
        if verdict is None:
            continue
        out.usable_lenient += 1
        # Absent means false: a schema-honouring model omitting a flag is
        # saying "not that", and treating a missing key as unknown would
        # flatter a model that simply answers less.
        said_clean = bool(verdict.get("matches"))
        if said_clean == case.truth.matches:
            out.correct_matches += 1
        if case.truth.matches and not said_clean:
            out.false_alarms += 1
        # A flag hit needs the right box ticked, not merely "something is
        # wrong": the loop picks its bounded fix off that flag, so a verdict
        # that is right for the wrong reason sends it down the wrong ladder.
        if not case.truth.matches and any(verdict.get(f) for f in case.truth.flags):
            out.flag_hits += 1
    return out


def report(cases: list[Case], scores: list[Score]) -> None:
    """Print the corpus provenance, then the arms, widest column first."""
    print("\n── corpus ──")
    for case in cases:
        how = case.truth.audit_kind or "by construction"
        note = "" if case.truth.visible else "   (renderer hides it — carried, not scored)"
        mark = "ok " if case.confirmed else "?? "
        print(f"  {mark}{case.name:16} confirmed by: {how}{note}")
    unconfirmed = [c.name for c in cases if not c.confirmed]
    if unconfirmed:
        print(f"  WARNING: unconfirmed labels: {', '.join(unconfirmed)}")

    scored = next((s.scored for s in scores), 0)
    print(f"\n── arms (n={scored} scored cases) ──")
    head = f"  {'model':18} {'mode':8} {'nores':>6} {'usable':>8} {'crash':>6} {'lenient':>8}"
    print(head + f" {'matches':>8} {'flags':>7} {'false':>6} {'median':>8}")
    for s in scores:
        # Everything after `nores` is out of `answered`, not out of `scored`:
        # a call the daemon never completed is not a judgement the model got
        # wrong, and dividing by it would quietly blame the model for it.
        print(
            f"  {s.model:18} {s.mode:8} {s.transport_failures:>6} "
            f"{s.usable_strict:>5}/{s.answered} {s.crashes:>6} "
            f"{s.usable_lenient:>5}/{s.answered} {s.correct_matches:>5}/{s.answered} "
            f"{s.flag_hits:>7} {s.false_alarms:>6} {s.median_s:>7.1f}s"
        )
    print(
        "\n  nores   = the daemon returned nothing (500/timeout) — not the model's fault,\n"
        "            and excluded from every denominator to its right\n"
        "  usable  = replies the current eyeball_slides consumer turns into a verdict\n"
        "  crash   = replies that consumer RAISES on (json.loads of prose)\n"
        "  lenient = replies a tolerant first-{...}-span parse recovers\n"
        "  matches = correct call on 'is this slide clean?' (lenient parse)\n"
        "  flags   = ...and ticked the right defect box\n"
        "  false   = flagged the clean control (each one costs a needless fix)"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument(
        "--models",
        default="gemma3:12b,qwen3-vl:4b",
        help="comma-separated Ollama tags to bench (default: %(default)s)",
    )
    parser.add_argument("--modes", default="strict,lenient", help="strict, lenient, or both")
    parser.add_argument("--json", type=Path, help="write the full results here")
    parser.add_argument("--keep", type=Path, help="keep the built deck and renders in this dir")
    args = parser.parse_args(argv)

    if not args.template.exists():
        print(f"template not found: {args.template}", file=sys.stderr)
        return 2

    with tempfile.TemporaryDirectory() as tmp:
        workdir = args.keep or Path(tmp)
        workdir.mkdir(parents=True, exist_ok=True)
        print(f"building corpus on {args.template.name} / {BENCH_LAYOUT} (page {BENCH_PAGE})…")
        cases, target = build_corpus(args.template, workdir)
        if args.keep:
            # Written only on request: looking at the pairs is how a reader
            # checks the bench measured what it says it measured.
            (workdir / "target.png").write_bytes(target)
            for case in cases:
                (workdir / f"cand_{case.name}.png").write_bytes(case.candidate)

        scores, raw_log = [], []
        for model in [m.strip() for m in args.models.split(",") if m.strip()]:
            for mode in [m.strip() for m in args.modes.split(",") if m.strip()]:
                print(f"\n  ── {model} / {mode} ──", flush=True)
                replies = run_arm(model, mode, cases, target)
                scores.append(score_arm(model, mode, cases, replies))
                raw_log += [
                    {
                        "model": model,
                        "mode": mode,
                        "case": r.case,
                        "seconds": round(r.seconds, 2),
                        "raw": r.raw if isinstance(r.raw, dict) else str(r.raw),
                        "strict_ok": r.strict is not None,
                        "crash": r.crash,
                        "lenient_ok": r.lenient is not None,
                    }
                    for r in replies
                ]
        report(cases, scores)
        if args.json:
            args.json.write_text(
                json.dumps(
                    {
                        "template": str(args.template),
                        "layout": BENCH_LAYOUT,
                        "page": BENCH_PAGE,
                        "cases": [
                            {"name": c.name, "confirmed": c.confirmed, "visible": c.truth.visible}
                            for c in cases
                        ],
                        "scores": [vars(s) for s in scores],
                        "replies": raw_log,
                    },
                    indent=2,
                ),
                encoding="utf-8",
            )
            print(f"\nwrote {args.json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
