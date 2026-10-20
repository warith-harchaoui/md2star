"""Seam-E assembler: clone the designer's own example slide, then replace its text.

:mod:`md2star.pptx_assemble` (Seam D) instantiates a *layout* and overlays
guaranteed-safe boxes in it. That is predictable, and it is also the ceiling of
what a layout can give you: a layout carries the background and the placeholder
boxes, and **nothing else**. The accent bars, the numbered step chevrons, the
decorative rules, the little bespoke shapes that make a deck look like *itself*
live on the designer's **example slides**, one level further in.

So this module takes the other half of the template. For a slide whose chosen
layout has an example page in the catalog (``LayoutInfo.example_pages``), it
**clones that page** and replaces its content: the designer's filler title
becomes yours, their filler bullets become yours, their stand-in picture slot
takes your image, and everything they drew around that content stays exactly
where they drew it. Typography is inherited rather than forced, which is the
whole point — a cloned paragraph keeps the designer's face, size, colour and
bullet glyph, so the result is their composition with your words in it.

Three rules keep that from being a liability.

**Nothing the designer wrote may ship.** This is the hard one, and it is the
same bug class as speaker notes printed on a slide: a template's filler text is
*their* copy ("Lorem ipsum", "Client name here"), and publishing it inside your
deck is worse than an empty box. So every text frame on a cloned slide is
either filled with your content or retired (:func:`_retire`) — there is no
third outcome — and a page whose content is not *restatable* at all (a chart, a
table, SmartArt, a group with words in it) is refused outright rather than
cloned with the designer's numbers still in it.

**A bad carrier is refused, not forced.** :func:`can_carry` decides, *before*
anything is cloned, whether an example page can hold this slide's content: a
three-column design needs at least three bullets, an image slide needs a
picture to replace. When it cannot, :func:`compose` returns ``None`` and the
caller falls back to Seam D, which can place anything. A predictable slide
beats a gutted composition.

**Legibility is still enforced.** Inheriting the designer's typography is the
feature; inheriting *unreadable* typography is not. Any run that declares a
size under :data:`md2star.pptx_audit.LEGIBLE_MIN_PT` is lifted to that floor. A
run that declares *no* size is left alone: it inherits from the designer's own
master, which is presentation-grade by construction — and the audit reads
declared sizes for the same reason.

Dependency direction is one-way, like the audit's: this module imports the
assembler's placeholder vocabulary and image placement, and
:func:`md2star.pptx_assemble.assemble` imports *this* one inside its own body,
so the pair never forms an import cycle.

Specified in ``docs/design/pptx-template-intelligence.md`` §10.4 (and §4.3 for
the A-D seams this one extends).

Author
------
[Warith HARCHAOUI](https://linkedin.com/in/warith-harchaoui/)
"""

from __future__ import annotations

import copy
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx.enum.text import MSO_AUTO_SIZE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt

from .logging import get_logger
from .pptx_assemble import (
    EMU_PER_INCH,
    LATENT_KINDS,
    TITLE_KINDS,
    add_image,
    is_ornament,
    placeholder_kind,
)
from .pptx_audit import LEGIBLE_MIN_PT
from .pptx_fit import AVG_CHAR_WIDTH_EM, EMU_PER_PT, FIT_SLACK, LINE_SPACING

logger = get_logger(__name__)

#: A text frame counts as a *peer* of the biggest one on the slide — one column
#: of a multi-column design, rather than a caption or a footnote — when it has
#: at least this share of the biggest one's area. Peers split the bullets
#: between them; everything else is retired.
PEER_AREA_RATIO = 0.4

#: A picture covering at least this share of the canvas is a *content* slot
#: (the designer's stand-in photo), so it is replaced by the slide's own image,
#: or removed with it when the slide has none. Anything smaller is furniture —
#: a logo, an icon, a rule — and stays exactly where the designer put it.
CONTENT_PICTURE_AREA = 0.04

#: Relationship ids are per-part, so every ``r:*`` attribute in the copied XML
#: has to be re-pointed at the clone's own relationship for the same target.
_REL_ATTRS = tuple(
    qn(f"r:{name}") for name in ("id", "embed", "link", "pict", "dm", "lo", "qs", "cs")
)

#: Relationships that must *not* be carried over: the layout is wired by
#: ``add_slide`` already, and a notes part belongs to exactly one slide —
#: pointing a second slide at the same one writes a file PowerPoint offers to
#: repair on open.
_SKIP_RELS = frozenset({RT.SLIDE_LAYOUT, RT.NOTES_SLIDE})

#: ``ppt/slides/slide12.xml`` -> 12. Example pages are numbered by *part*, so
#: this is how a catalog page number finds its slide (see
#: :func:`example_slides`).
_SLIDE_PART = re.compile(r"slide(\d+)\.xml\Z")


@dataclass
class Content:
    """What a Markdown slide wants placed, independent of how it gets placed.

    The assembler has made every content decision by the time Seam E is asked —
    which branch the slide took, how many bullets survived the eyeball loop's
    cap, where a Mermaid diagram was rendered to — so this carries the outcome
    and both seams consume the same resolved payload.
    """

    title: str
    bullets: list[str] = field(default_factory=list)
    image: Path | None = None
    alt: str = ""


@dataclass
class Roles:
    """Every shape of a slide, sorted into what Seam E may do with it."""

    #: Where the slide's heading goes: the designer's own title placeholder
    #: when the page has one, else its largest, highest type.
    title: object | None = None
    #: Text frames that can carry bullets, in reading order
    #: (:func:`reading_order`: columns left to right, rows top to bottom).
    bodies: list = field(default_factory=list)
    #: Pictures big enough to be content slots, biggest first.
    pictures: list = field(default_factory=list)
    #: Design, not content: decorative shapes, small pictures, and the latent
    #: date / footer / slide-number placeholders. Never touched.
    furniture: list = field(default_factory=list)
    #: Content Seam E cannot restate — a chart, a table, SmartArt, a group with
    #: words in it. One of these is enough to disqualify the whole page, since
    #: cloning it would publish the designer's own numbers (see
    #: :func:`can_carry`).
    foreign: list = field(default_factory=list)


# ── reading the template's example slides ────────────────────────────────────
def example_slides(prs) -> dict[int, object]:
    """Map each example *page number* to the slide object that is that page.

    Keyed by the number in the slide's part name rather than by position in the
    deck, because that is how :func:`md2star.pptx_layout.layout_pages` numbers
    the pages it records in ``LayoutInfo.example_pages``. The two agree even in
    a template whose slides have been reordered in PowerPoint since they were
    created, where part order and presentation order no longer match.

    Snapshot it *before* cloning anything: the clones are slides of the same
    presentation, and a caller that re-read this map mid-build would start
    finding them in it.
    """
    pages: dict[int, object] = {}
    for slide in prs.slides:
        match = _SLIDE_PART.search(str(slide.part.partname))
        if match:
            pages[int(match.group(1))] = slide
    return pages


def template_slide_ids(prs) -> list:
    """The ``<p:sldId>`` entries the template shipped with.

    Seam E has to *read* the designer's example slides while it builds, so
    unlike Seam D it cannot clear them up front
    (:func:`md2star.pptx_assemble.clear_slides`). It records them here, appends
    its own slides after them, and drops these at the end — see
    :func:`drop_template_slides`. Authored order survives either way, since
    both seams append.
    """
    return list(prs.slides._sldIdLst)  # noqa: SLF001 — no public remove-slide API


def drop_template_slides(prs, ids: list) -> None:
    """Remove the example slides *ids*, leaving only what the build appended.

    The images on those pages are not dropped with them when a clone uses one:
    :func:`_copy_rels` gives every clone its *own* relationship to the same
    image part, so the part stays referenced and gets written.
    """
    lst = prs.slides._sldIdLst  # noqa: SLF001 — same workaround as clear_slides
    for sld in ids:
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


# ── classifying a page's shapes ──────────────────────────────────────────────
def _area(shape) -> int:
    """Shape area in EMU². ``0`` when the geometry is not resolvable."""
    return (shape.width or 0) * (shape.height or 0)


def _is_picture(shape) -> bool:
    # The element tag, not ``shape_type``: a picture *placeholder* the designer
    # filled is also a <p:pic>, and that is exactly the case we must catch.
    return shape.element.tag == qn("p:pic")


def _is_textbox(shape) -> bool:
    """True for a bare text box, false for an autoshape or a placeholder.

    OOXML records the difference itself (``<p:cNvSpPr txBox="1">``), which is
    better than guessing from the fill: a text box exists only for its text, so
    an unused one is filler and can go, while an unused *autoshape* is a drawn
    object — a chevron, a card, a band — that the composition needs.
    """
    properties = shape.element.find(f".//{qn('p:cNvSpPr')}")
    return properties is not None and properties.get("txBox") == "1"


def _deep_text(shape) -> str:
    """All text anywhere under *shape*, including inside a group."""
    return "".join(node.text or "" for node in shape.element.iter(qn("a:t"))).strip()


def _max_pt(shape) -> float:
    """Largest *declared* run size in points (``0.0`` when all are inherited)."""
    if not shape.has_text_frame:
        return 0.0
    sizes = [
        run.font.size.pt
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    return max(sizes, default=0.0)


def _overlaps_horizontally(a, b) -> bool:
    """Whether two frames occupy the same vertical band of the slide."""
    left_a, left_b = (a.left or 0), (b.left or 0)
    right_a, right_b = left_a + (a.width or 0), left_b + (b.width or 0)
    shared = min(right_a, right_b) - max(left_a, left_b)
    narrower = min((a.width or 0), (b.width or 0))
    return narrower > 0 and shared > 0.5 * narrower


def reading_order(shapes: list) -> list:
    """Order frames the way a reader walks them: columns first, then rows.

    Sorting by ``(top, left)`` is right for a row of columns and wrong for a
    split design, where it interleaves the two halves — the left panel's one
    line landing between the right panel's two. Found on a real template, not
    reasoned about: a cloned Jellysmack split page took bullet 2 into its lime
    panel and bullets 1 and 3 into the black half.

    So frames are grouped into **columns** (overlapping vertical bands), the
    columns run left to right, and each column runs top to bottom. A single row
    of boxes degenerates to left-to-right, a single stack to top-to-bottom, and
    the split design comes out in the order a reader actually reads it.
    """
    columns: list[list] = []
    for shape in sorted(shapes, key=lambda s: s.left or 0):
        for column in columns:
            if _overlaps_horizontally(column[0], shape):
                column.append(shape)
                break
        else:
            columns.append([shape])
    return [shape for column in columns for shape in sorted(column, key=lambda s: s.top or 0)]


def classify(slide, canvas: tuple[int, int]) -> Roles:
    """Sort every shape of *slide* into a :class:`Roles`.

    Parameters
    ----------
    slide:
        The page being classified — a designer example, or the clone of one.
    canvas:
        ``(width, height)`` of the deck in EMU; a picture's share of it is what
        separates a content slot from a logo.
    """
    canvas_area = max(1, canvas[0] * canvas[1])
    roles = Roles()
    texts: list = []
    for shape in slide.shapes:
        if placeholder_kind(shape) in LATENT_KINDS:
            # Date / footer / slide number: empty is their correct state,
            # PowerPoint fills them itself.
            roles.furniture.append(shape)
        elif shape.element.tag == qn("p:graphicFrame"):
            # A chart, a table or SmartArt carries the designer's own data and
            # there is no "replace the text" that restates it.
            roles.foreign.append(shape)
        elif shape.element.tag == qn("p:grpSp"):
            # A group is design when it is shapes (an icon, a bracket, a badge)
            # and content when it has words in it — and its words are out of
            # reach of the retire-or-fill rule, since ``slide.shapes`` does not
            # descend into groups.
            design = is_ornament(_deep_text(shape))
            (roles.furniture if design else roles.foreign).append(shape)
        elif _is_picture(shape):
            slot = _area(shape) >= CONTENT_PICTURE_AREA * canvas_area
            (roles.pictures if slot else roles.furniture).append(shape)
        elif shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape)
        else:
            roles.furniture.append(shape)

    roles.pictures.sort(key=lambda shape: -_area(shape))
    named_titles = [shape for shape in texts if placeholder_kind(shape) in TITLE_KINDS]
    if named_titles:
        # The designer's own answer, and the one that feeds the outline pane
        # and a screen reader's reading order.
        roles.title = named_titles[0]
    elif texts:
        # No title placeholder on the page: the biggest declared type wins,
        # ties broken by whichever sits higher — a heading is big and near the
        # top, and on a page where nothing declares a size, the top one is
        # still the better guess than the first in XML order.
        roles.title = max(texts, key=lambda s: (_max_pt(s), -(s.top or 0), _area(s)))
    roles.bodies = reading_order([shape for shape in texts if shape is not roles.title])
    return roles


def fits(shape, lines: list[str]) -> bool:
    """Whether *lines* can sit in *shape* without dropping under the floor.

    The designer sized that frame for the designer's own line. Ours may be
    three times as long, and the renderer answers overflow by shrinking the
    type — quietly, which is the point: a title that falls from 38 pt to 27 pt
    is still a title, while three bullets squeezed into a one-line kicker come
    out at 9 pt and nobody in the room can read them.

    So the question is not "does it overflow" but "how small would it have to
    get". Same character model as the rest of the suite (0.5 em mean advance,
    1.2 line spacing) and the same square-root implied-size arithmetic as
    :func:`md2star.pptx_audit._check_overrun` — capacity scales with the square
    of the point size, since line height *and* characters per line both do.
    A frame that declares no size at all is left alone: its typography comes
    from the designer's master, and no signal is not a bad signal.
    """
    declared = _max_pt(shape)
    available = (shape.height or 0) / EMU_PER_PT
    width_pt = (shape.width or 0) / EMU_PER_PT
    if declared <= 0 or available <= 0 or width_pt <= 0:
        return True
    chars_per_line = max(1.0, width_pt // (declared * AVG_CHAR_WIDTH_EM))
    needed = sum(math.ceil(len(line) / chars_per_line) for line in lines) * declared * LINE_SPACING
    if needed <= available * FIT_SLACK:
        return True
    return declared * math.sqrt(available / needed) >= LEGIBLE_MIN_PT


def peers(bodies: list) -> list:
    """The body frames that are columns of one design, not captions of it."""
    if not bodies:
        return []
    biggest = max(_area(shape) for shape in bodies)
    if biggest <= 0:
        return list(bodies)  # no resolvable geometry: treat them all as equals
    return [shape for shape in bodies if _area(shape) >= PEER_AREA_RATIO * biggest]


def can_carry(roles: Roles, content: Content) -> bool:
    """Whether a page classified as *roles* can actually hold *content*.

    The refusals are the interesting part, and each one is a composition this
    seam would otherwise break:

    * **no title frame** — nowhere to put the heading, and a cloned page with
      the designer's heading still on it is the leak this module exists to
      avoid;
    * **content that cannot be restated** (``roles.foreign``) — a chart or a
      table would ship with the designer's numbers in it;
    * **an image with no slot to put it in** — the picture would land on top of
      a composition that was not drawn around it;
    * **a slot with no image to put in it** — the mirror case, and the one a
      vision model catches first: a page the designer drew around a photograph
      keeps the shape of that photograph, so taking it out (which we must, it
      is their content) leaves the void the rest of the page was balanced
      against. A text slide belongs on a page the designer built for text;
    * **fewer bullets than columns** — a three-column design with one filled
      column reads as broken;
    * **text that would have to shrink past the legibility floor** to fit the
      frames the designer sized for their own line (:func:`fits`) — the page
      that taught this one put three bullets in a kicker meant for four words,
      and the renderer dutifully set them at 9 pt.

    Every refusal is the same move: the page is declined and Seam D, which can
    place anything, takes the slide instead.
    """
    if roles.title is None or roles.foreign:
        return False
    if (content.image is not None) != bool(roles.pictures):
        return False
    if content.bullets:
        columns = peers(roles.bodies)
        if not columns or len(content.bullets) < len(columns):
            return False
        shares = distribute(content.bullets, len(columns))
        return all(fits(frame, share) for frame, share in zip(columns, shares, strict=True))
    return True


def distribute(bullets: list[str], frames: int) -> list[list[str]]:
    """Split *bullets* across *frames* columns in reading order.

    Near-equal contiguous chunks, so a three-column design stays balanced
    instead of overflowing its first column and leaving the other two thin.
    :func:`can_carry` guarantees there are at least as many bullets as frames,
    so no column comes out empty.
    """
    if frames <= 1:
        return [list(bullets)]
    shares: list[list[str]] = []
    start = 0
    for position in range(frames):
        take = len(bullets) // frames + (1 if position < len(bullets) % frames else 0)
        shares.append(list(bullets[start : start + take]))
        start += take
    return shares


# ── cloning a page ───────────────────────────────────────────────────────────
def _copy_rels(src_part, dst_part) -> dict[str, str]:
    """Relate *dst_part* to everything *src_part* points at; old rId -> new."""
    mapping: dict[str, str] = {}
    for rel in src_part.rels.values():
        if rel.reltype in _SKIP_RELS:
            continue
        if rel.is_external:
            mapping[rel.rId] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            # The same part object, not a copy of it: two slides sharing one
            # image part is how PowerPoint itself stores a reused photo.
            mapping[rel.rId] = dst_part.relate_to(rel.target_part, rel.reltype)
    return mapping


def _retarget(element, mapping: dict[str, str]) -> None:
    """Rewrite every relationship reference under *element* through *mapping*."""
    for node in element.iter():
        for attribute in _REL_ATTRS:
            old = node.get(attribute)
            if old is None:
                continue
            if old in mapping:
                node.set(attribute, mapping[old])
            else:
                # Only the layout and notes relationships are deliberately not
                # copied, and neither is referenced from a shape — so this is
                # worth a line in a verbose log rather than silence.
                logger.debug("clone: %s=%r has no copied relationship", attribute, old)


def clone_slide(prs, src):
    """Duplicate *src* — a slide of *prs* — and return the new slide.

    python-pptx has no slide-copy API; this is the shape-tree deepcopy plus
    relationship re-pointing that stands in for one. It stays *inside* one
    presentation on purpose: the deck Seam E builds is opened from the designer
    template itself, so their example pages are already in the same package and
    no cross-file part surgery is needed.
    """
    dst = prs.slides.add_slide(src.slide_layout)
    # ``add_slide`` clones the LAYOUT's placeholders onto the new slide. The
    # example page already carries its own filled copies of them, so leaving
    # ours in place would ship two shapes per region — the overlay/placeholder
    # defect §10.1 closed for Seam D, walked back in from the other side.
    for shape in list(dst.shapes):
        shape.element.getparent().remove(shape.element)

    rels = _copy_rels(src.part, dst.part)
    tree = dst.shapes._spTree  # noqa: SLF001 — python-pptx exposes no shape-insert API
    for shape in src.shapes:
        element = copy.deepcopy(shape.element)
        _retarget(element, rels)
        tree.append(element)

    # A slide-level background override (a photo, a gradient the layout itself
    # does not carry) is part of the composition and lives outside the shape
    # tree, so a shape-only copy would silently lose it.
    src_csld = src.element.find(qn("p:cSld"))
    dst_csld = dst.element.find(qn("p:cSld"))
    background = None if src_csld is None else src_csld.find(qn("p:bg"))
    if background is not None and dst_csld is not None:
        copied = copy.deepcopy(background)
        _retarget(copied, rels)
        dst_csld.insert(0, copied)  # <p:bg> must precede <p:spTree>
    return dst


# ── replacing a page's content ───────────────────────────────────────────────
def _retext(paragraph, text: str) -> None:
    """Make *paragraph* say exactly *text*, in its first run's formatting.

    Run-level ``rPr`` is what carries the designer's face, size, weight and
    theme colour, and nothing in python-pptx's public API copies it — so the
    run is *kept* and only its text replaced. Everything after it goes: a
    second run, a line break, a field (their slide number, their date) are all
    the designer's own, and a cloned paragraph must say one thing, ours.
    """
    runs = paragraph.findall(qn("a:r"))
    for extra in runs[1:]:
        paragraph.remove(extra)
    for tag in ("a:br", "a:fld"):
        for node in paragraph.findall(qn(tag)):
            paragraph.remove(node)
    run = runs[0] if runs else paragraph.add_r()
    # A hyperlink lives in the run's own ``rPr``, so it would survive the text
    # swap and leave our words carrying a link the author never wrote — usually
    # to a page of the template that is not even in this deck. It is also the
    # only shape-level reference a slide can hold to *another slide*, so
    # dropping it is what keeps a cloned page from pinning an example page the
    # build is about to remove.
    properties = run.find(qn("a:rPr"))
    if properties is not None:
        for tag in ("a:hlinkClick", "a:hlinkMouseOver"):
            for link in properties.findall(qn(tag)):
                properties.remove(link)
    run.text = text


def _enforce_floor(shape, font_delta: int = 0) -> None:
    """Apply *font_delta*, never going under the legibility floor.

    A run with no declared size is left untouched: its size comes from the
    designer's own master, which is the typography this seam exists to inherit,
    and :mod:`md2star.pptx_audit` reads declared sizes for the same reason.
    """
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is None:
                continue
            run.font.size = Pt(max(LEGIBLE_MIN_PT, run.font.size.pt + font_delta))


def _fill_frame(shape, lines: list[str], *, font_delta: int = 0) -> None:
    """Replace a designer frame's paragraphs with one per line of *lines*.

    Every line is a copy of the designer's own first paragraph, so indent
    level, bullet glyph and run typography come along with it — which is why
    Seam E does *not* prefix a "• " the way Seam D must: the master already
    draws the glyph, and adding ours would double it.
    """
    # A text frame with no paragraph at all is invalid OOXML, so a frame asked
    # for nothing still gets one — empty. Unreachable through `compose`, whose
    # feasibility check guarantees a line per column, and cheaper to hold than
    # to prove impossible for every future caller.
    lines = lines or [""]
    frame = shape.text_frame
    body = frame._txBody  # noqa: SLF001 — paragraph copying needs the element
    paragraphs = body.findall(qn("a:p"))
    # Prefer a paragraph that actually has a run to inherit from; a designer's
    # first line is occasionally an empty spacer.
    with_runs = [p for p in paragraphs if p.findall(qn("a:r"))]
    prototype = copy.deepcopy((with_runs or paragraphs)[0])
    for paragraph in paragraphs:
        body.remove(paragraph)
    for line in lines:
        clone = copy.deepcopy(prototype)
        _retext(clone, line)
        body.append(clone)

    frame.word_wrap = True
    # Belt and braces, exactly as Seam D: our line is not the designer's filler
    # and may be longer, so the renderer shrinks it rather than clipping it.
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    _enforce_floor(shape, font_delta)


def _retire(shape) -> None:
    """Guarantee a frame the deck did not fill ships nothing the designer wrote.

    Three outcomes, in order: ornament stays (the "01" of a numbered step is
    the design, not filler), a bare text box is removed (it existed only for
    the words we are taking out of it), and an autoshape or placeholder is
    emptied but kept (it is a drawn object; the assembler's own sweep then
    removes any placeholder left empty).
    """
    if is_ornament(shape.text_frame.text.strip()):
        return
    if _is_textbox(shape):
        shape.element.getparent().remove(shape.element)
        return
    shape.text_frame.clear()


def _replace_picture(slide, picture, image: Path, alt: str) -> None:
    """Swap the designer's stand-in photo for *image*, inside its own slot.

    Their frame is kept as the *box*, not as the frame: contain-fitting into it
    (:func:`md2star.pptx_assemble.add_image`) preserves the composition's slot
    while refusing to stretch a picture whose aspect does not match the
    stand-in's. A cropped stand-in's crop goes with it, which is the price of
    never distorting the user's own image.
    """
    box = (
        (picture.left or 0) / EMU_PER_INCH,
        (picture.top or 0) / EMU_PER_INCH,
        (picture.width or 0) / EMU_PER_INCH,
        (picture.height or 0) / EMU_PER_INCH,
    )
    element = picture.element
    parent = element.getparent()
    position = list(parent).index(element)
    parent.remove(element)
    add_image(slide, image, box=box, alt=alt)
    # Stacking order is composition: the designer may have text over this
    # photo, and ``add_picture`` appends, which would bury it. Put the new
    # picture back exactly where the old one sat in the shape tree.
    placed = list(parent)[-1]
    parent.remove(placed)
    parent.insert(position, placed)


# ── the seam ─────────────────────────────────────────────────────────────────
def _drop_unused_rels(slide) -> int:
    """Drop the clone's relationships nothing on it points at any more.

    :func:`_copy_rels` has to relate the clone to *everything* the example page
    pointed at, because the copied XML is re-targeted afterwards. Whatever the
    content replacement then removed — the designer's stand-in photo, most of
    all — leaves its relationship behind, and a relationship is what keeps a
    part in the saved file. Without this sweep their stock photo ships inside
    the deck, invisible but present, which is the same mistake as their filler
    text, only harder to notice.
    """
    referenced = set()
    for node in slide.element.iter():
        for attribute in _REL_ATTRS:
            value = node.get(attribute)
            if value is not None:
                referenced.add(value)
    dropped = 0
    for rId, rel in list(slide.part.rels.items()):
        # The layout relationship is never referenced from the slide's own XML
        # and is mandatory; a notes relationship does not exist yet (the
        # assembler adds the notes after this). Both are in _SKIP_RELS.
        if rel.reltype in _SKIP_RELS or rId in referenced:
            continue
        slide.part.drop_rel(rId)
        dropped += 1
    return dropped


def pick_example(
    examples: dict[int, object],
    pages: list[int],
    content: Content,
    *,
    used: set[int],
    canvas: tuple[int, int],
    avoid: list[int] | tuple[int, ...] = (),
) -> tuple[int, object] | None:
    """Choose which example page to clone for *content*, or ``None``.

    Candidates are the chosen layout's own ``example_pages``, in the designer's
    order, narrowed to those that :func:`can_carry` the content. Among those, a
    page this deck has not cloned yet wins — so a layout demonstrated by three
    pages gives three different-looking slides instead of the same one three
    times — and otherwise the first feasible page is reused. A slide with an
    image then takes the page offering the **biggest** picture slot, since
    contain-fitting means that slot is the size the reader actually sees.

    *avoid* is the hard exclusion the Ralph Eyeball Loop writes: a page a
    vision model has already judged not to be carrying this slide. *used* is
    only a preference (variety across the deck); *avoid* is a refusal. When
    every feasible page has been refused, the answer is ``None`` and the slide
    falls back to Seam D, which is the safe generic the loop wants anyway.
    """
    feasible: list[tuple[int, object, int]] = []
    for page in pages:
        if page in avoid:
            continue
        slide = examples.get(page)
        if slide is None:
            continue
        roles = classify(slide, canvas)
        if can_carry(roles, content):
            # How much room this page gives a picture, which is the one thing
            # worth ranking the designer's own order by: a chart dropped into a
            # thumbnail-sized slot is the most visible way this seam can fail,
            # and contain-fitting means the slot is exactly what the reader
            # ends up seeing. Zero for a text slide, so their order stands.
            room = _area(roles.pictures[0]) if (content.image and roles.pictures) else 0
            feasible.append((page, slide, room))
    if not feasible:
        return None
    fresh = [c for c in feasible if c[0] not in used] or feasible
    page, slide, _ = max(fresh, key=lambda c: c[2]) if content.image else fresh[0]
    return page, slide


def compose(prs, example, content: Content, *, font_delta: int = 0):
    """Clone *example* and put *content* in it; ``None`` if it cannot carry it.

    Classification runs twice on purpose: once on *example* to decide
    feasibility before anything is written, and once on the clone, whose shapes
    are different objects. Deciding first is what makes a refusal free of side
    effects — the caller gets ``None`` and a presentation it can still build a
    Seam-D slide on.

    Returns
    -------
    object | None
        The new slide, or ``None`` when the page was refused.
    """
    canvas = (prs.slide_width, prs.slide_height)
    if not can_carry(classify(example, canvas), content):
        return None

    slide = clone_slide(prs, example)
    roles = classify(slide, canvas)
    if roles.title is not None:
        _fill_frame(roles.title, [content.title], font_delta=font_delta)

    columns: list = []
    if content.bullets:
        columns = peers(roles.bodies)
        for frame, share in zip(columns, distribute(content.bullets, len(columns)), strict=True):
            _fill_frame(frame, share, font_delta=font_delta)
    for frame in roles.bodies:
        if frame not in columns:
            _retire(frame)

    spare = roles.pictures
    if content.image is not None and roles.pictures:
        _replace_picture(slide, roles.pictures[0], content.image, content.alt)
        spare = roles.pictures[1:]
    for picture in spare:
        # A content slot with nothing of ours to put in it: the designer's
        # stand-in photo is their content, and shipping it inside someone
        # else's deck is the same mistake as shipping their filler text.
        picture.element.getparent().remove(picture.element)

    orphans = _drop_unused_rels(slide)
    if orphans:
        logger.debug("clone: dropped %d relationship(s) nothing points at", orphans)
    return slide
